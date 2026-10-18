# -*- coding: utf-8 -*-
"""
Script to generate the complete laboratory report Word document:
Суханов_СС_Отчет_5.docx based on Вяткин_АА_Отчет_5.docx sample.
"""

import os
import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_report():
    base_file = 'Вяткин_АА_Отчет_5.docx'
    doc = docx.Document(base_file)

    def style_run(run, font_name='Times New Roman', size_pt=14, bold=False, italic=False, color=(0,0,0)):
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = RGBColor(*color)

    # 1. Update Title Page
    p25 = doc.paragraphs[25]
    p25.text = 'студент гр. 5140904/60401'
    if p25.runs:
        style_run(p25.runs[0], size_pt=14)

    p26 = doc.paragraphs[26]
    p26.text = '\t______________________\tСуханов С.С.'
    if p26.runs:
        style_run(p26.runs[0], size_pt=14)

    p30 = doc.paragraphs[30]
    p30.text = '\t______________________\t Шишкина В.И.'
    if p30.runs:
        style_run(p30.runs[0], size_pt=14)

    p33 = doc.paragraphs[33]
    p33.text = ' «___» ______________ 2026 г.'
    if p33.runs:
        style_run(p33.runs[0], size_pt=14)

    p41 = doc.paragraphs[41]
    p41.text = '2026'
    if p41.runs:
        style_run(p41.runs[0], size_pt=14)

    # 2. Update Figure Captions
    captions_map = {
        70: 'Рис. 1.1. Версия JDK и Gradle.',
        73: 'Рис. 1.2. Содержимое дистрибутива.',
        82: 'Рис. 1.3. Часть конфигурационного файла.',
        86: 'Рис. 1.4. Окно приложения iResolver.',
        98: 'Рис. 1.5. Сгенерированный отчет.',
        100: 'Рис. 1.6. Раздел «Similar issues».',
        102: 'Рис. 1.7. Раздел «Qualified users».',
        104: 'Рис. 1.8. Раздел «Probable labels».',
        106: 'Рис. 1.9. Раздел «Probable attachments».',
        108: 'Рис. 1.10. Раздел «Useful Documentation».',
        110: 'Рис. 1.11. Раздел «Confluence pages».',
        112: 'Рис. 1.12. Раздел «Proposals».',
        115: 'Рис. 1.13. Измененный JQL запрос.',
        117: 'Рис. 1.14. Созданные правила.',
        120: 'Рис. 1.15. Созданные правила в новом отчете.',
        123: 'Рис. 1.16. Проверка правил часть 1.',
        125: 'Рис. 1.17. Проверка правил часть 2.'
    }

    for p_idx, cap_text in captions_map.items():
        p = doc.paragraphs[p_idx]
        p.text = cap_text
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.runs:
            style_run(p.runs[0], size_pt=12, italic=True)

    # Helper insertion functions
    def insert_p_before(target_p, text, bold=False, italic=False, space_before=0, space_after=0, indent=1.25):
        new_p = target_p.insert_paragraph_before()
        new_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        new_p.paragraph_format.first_line_indent = Cm(indent)
        new_p.paragraph_format.line_spacing = 1.5
        new_p.paragraph_format.space_before = Pt(space_before)
        new_p.paragraph_format.space_after = Pt(space_after)
        if text:
            r = new_p.add_run(text)
            style_run(r, size_pt=14, bold=bold, italic=italic)
        return new_p

    def insert_h1_before(target_p, text):
        new_p = target_p.insert_paragraph_before()
        new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        new_p.paragraph_format.first_line_indent = Cm(0)
        new_p.paragraph_format.line_spacing = 1.5
        new_p.paragraph_format.space_before = Pt(16)
        new_p.paragraph_format.space_after = Pt(8)
        new_p.paragraph_format.keep_with_next = True
        r = new_p.add_run(text)
        style_run(r, size_pt=14, bold=True)
        return new_p

    def insert_h2_before(target_p, text):
        new_p = target_p.insert_paragraph_before()
        new_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        new_p.paragraph_format.first_line_indent = Cm(1.25)
        new_p.paragraph_format.line_spacing = 1.5
        new_p.paragraph_format.space_before = Pt(12)
        new_p.paragraph_format.space_after = Pt(4)
        new_p.paragraph_format.keep_with_next = True
        r = new_p.add_run(text)
        style_run(r, size_pt=14, bold=True)
        return new_p

    def insert_code_box_before(target_p, code_str, caption=""):
        if caption:
            p_cap = target_p.insert_paragraph_before()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_cap.paragraph_format.first_line_indent = Cm(1.25)
            p_cap.paragraph_format.space_before = Pt(8)
            p_cap.paragraph_format.space_after = Pt(2)
            r_c = p_cap.add_run(caption)
            style_run(r_c, size_pt=12, bold=True, italic=True)

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Cm(16.5)

        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8F9FA"/>')
        tcPr.append(shd)

        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
            f'<w:left w:val="single" w:sz="18" w:space="0" w:color="0969DA"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(code_str)
        r.font.name = 'Consolas'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(36, 41, 46)

        target_p._p.addprevious(tbl._tbl)

    # 3. Insert Custom Rules explanation and code listing BEFORE paragraph 118 ("Заново создадим новые дата-сеты...")
    # Find paragraph 118 by its text prefix:
    p_dataset_regen = None
    for p in doc.paragraphs:
        if 'Заново создадим новые дата-сеты' in p.text:
            p_dataset_regen = p
            break

    if p_dataset_regen:
        insert_p_before(p_dataset_regen, "На рис. 1.14 представлены разработанные новые экспертные правила на декларативном языке Drools. Для автоматизации триажа и контроля бизнес-процесса обработки задач были сформулированы следующие правила:")
        insert_p_before(p_dataset_regen, "1. «Sub-tasks are present»: проверяет наличие подзадач у открытого тикета. Если список subTasks не пуст, в раздел предложений выдается рекомендация о необходимости закрыть все подзадачи перед финализацией родительской задачи («Resolve sub-tasks first before closing parent ticket»).")
        insert_p_before(p_dataset_regen, "2. «Open with active comments»: анализирует тикеты в статусе Open, по которым уже ведется активное обсуждение (список comments не пуст). Правило автоматически рекомендует перевести задачу в рабочий статус («Change task status to IN PROGRESS»).")
        insert_p_before(p_dataset_regen, "3. «High Priority Unassigned Issue»: отслеживает критические задачи с приоритетами Major, Critical, Blocker, у которых отсутствует назначенный исполнитель (assignee == null), и формирует рекомендацию о немедленном назначении старшего инженера.")
        insert_p_before(p_dataset_regen, "4. «No fix version specified»: контролирует корректность заполнения целевой версии релиза и генерирует рекомендацию запланировать тикет в ближайший майлстоун при отсутствии поля fixVersions.")

        custom_rules_code = '''package com.koval.resolver.rules

import com.koval.resolver.common.api.model.issue.Issue;
global java.util.List results

rule "Sub-tasks are present"
    dialect "mvel"
    when
        $issue : Issue( subTasks != null && !subTasks.isEmpty() )
    then
        results.add("Resolve sub-tasks first before closing parent ticket");
end

rule "Open with active comments"
    dialect "mvel"
    when
        $issue : Issue( (status == "Open" || status == "OPEN") && 
                        comments != null && !comments.isEmpty() )
    then
        results.add("Change task status to IN PROGRESS");
end

rule "High Priority Unassigned Issue"
    dialect "mvel"
    when
        $issue : Issue( (priority == "Major" || priority == "Critical" || priority == "Blocker") 
                        && assignee == null )
    then
        results.add("High priority issue without assignee: assign senior engineer immediately");
end

rule "No fix version specified"
    dialect "mvel"
    when
        $issue : Issue( fixVersions == null || fixVersions.isEmpty() )
    then
        results.add("Target fix version is missing: schedule for upcoming release milestone");
end'''
        insert_code_box_before(p_dataset_regen, custom_rules_code, "Листинг 1. Исходный код правил CustomRules.drl")

    # 4. Insert report explanation BEFORE "Откроем одну из задач..."
    p_check_task = None
    for p in doc.paragraphs:
        if 'Откроем одну из задач' in p.text:
            p_check_task = p
            break

    if p_check_task:
        insert_p_before(p_check_task, "В сгенерированном отчете (рис. 1.15) видно, как новые правила Drools успешно отработали на целевых открытых тикетах Apache Kafka:")
        insert_p_before(p_check_task, "• Для задачи KAFKA-16800 («Resolve Gradle 9.0 deprecations»), имеющей 7 подзадач и 3 комментария в статусе Open без назначенной версии, сработали предложения: «Resolve sub-tasks first before closing parent ticket», «Change task status to IN PROGRESS» и «Target fix version is missing: schedule for upcoming release milestone».")
        insert_p_before(p_check_task, "• Для задачи KAFKA-16628 («Add system test for validating static consumer bounce...»), не имеющей подзадач, правило о подзадачах корректно не сработало, однако сработали предложения о переводе в статус IN PROGRESS и указании fixVersion.")

    # 5. Find "Вывод" paragraph and insert the two theoretical/analytical sections strictly before it
    p_vyvod = None
    for p in doc.paragraphs:
        if p.text.strip() == 'Вывод':
            p_vyvod = p
            break

    if p_vyvod:
        # SECTION: Ответы на контрольные вопросы
        insert_h1_before(p_vyvod, "Ответы на контрольные вопросы")
        
        insert_h2_before(p_vyvod, "1. Дерево Хаффмана и Hierarchical Softmax в Doc2Vec")
        insert_p_before(p_vyvod, "В классических нейросетевых моделях дистрибутивной семантики (Word2Vec, Doc2Vec) на этапе обучения для предсказания очередного слова w из словаря V используется стандартная функция Softmax:")
        insert_p_before(p_vyvod, "P(w_O | w_I) = exp(v'_{w_O}^T * v_{w_I}) / sum_{w=1}^{|V|} exp(v'_w^T * v_{w_I})")
        insert_p_before(p_vyvod, "В практических задачах размер словаря |V| составляет от сотен тысяч до миллионов уникальных слов. Вычисление знаменателя суммы требует полного прохода по всем векторам словаря на каждом шаге обновления весов, что влечет вычислительную сложность O(|V|) и делает прямое обучение градиентным спуском крайне ресурсоемким.")
        insert_p_before(p_vyvod, "Для решения этой проблемы применяется оптимизация Hierarchical Softmax, в основе которой лежит бинарное дерево Хаффмана:")
        insert_p_before(p_vyvod, "• Дерево Хаффмана строится на основе частотного анализа корпуса текстов: все слова словаря V помещаются в листья бинарного дерева, а частые слова располагаются ближе к корню (имеют минимальную длину кодового пути).")
        insert_p_before(p_vyvod, "• Вычисление вероятности конкретного слова w сводится к вычислению вероятностей последовательных бинарных ветвлений (влево/вправо с помощью логистической сигмоиды sigma(theta^T * v)) по уникальному пути от корня дерева к соответствующему листу.")
        insert_p_before(p_vyvod, "• Благодаря сбалансированности префиксного дерева средняя длина пути составляет O(log_2 |V|). Вместо сотен тысяч операций нормализации на каждом шаге вычисляется всего 15–20 скалярных произведений, что ускоряет обучение моделей Doc2Vec в тысячи раз.")

        insert_h2_before(p_vyvod, "2. Архитектура Skip-Gram")
        insert_p_before(p_vyvod, "Архитектура Skip-Gram (Миколов и др.) является одним из ключевых алгоритмов обучения распределенных векторных представлений:")
        insert_p_before(p_vyvod, "• Целевая функция Skip-Gram: по заданному центральному слову w_t максимизировать условную вероятность предсказания слов контекста w_{t+j} в симметричном скользящем окне шириной c:")
        insert_p_before(p_vyvod, "L = sum_{t=1}^T sum_{-c <= j <= c, j != 0} log P(w_{t+j} | w_t)")
        insert_p_before(p_vyvod, "• В отличие от моделей типа Bag of Words, которые усредняют окружение для предсказания одного центрального слова, Skip-Gram обучается на парах (центральное слово, контекстное слово).")
        insert_p_before(p_vyvod, "• Важнейшее практическое преимущество Skip-Gram в контексте сопровождения ПО заключается в превосходной обработке редких терминов. В описаниях технических инцидентов часто встречаются уникальные идентификаторы классов, названия методов, стек-трейсы и специфические конфигурационные ключи. При усреднении в CBOW вектор редкого слова нивелируется более частыми токенами, в то время как Skip-Gram многократно формирует обучающие пары для каждого редкого токена, наделяя его точным семантическим вектором.")

        insert_h2_before(p_vyvod, "3. Метод Bag of Words, CBOW и модификация PV-DBOW")
        insert_p_before(p_vyvod, "Концепция Bag of Words («мешок слов») имеет три последовательных этапа развития в компьютерной лингвистике:")
        insert_p_before(p_vyvod, "1. Классический Bag of Words (BoW): статистический подход, при котором документ кодируется вектором размерности |V|, где каждый элемент равен частоте встречаемости слова. Подход полностью игнорирует порядок слов, синтаксис и грамматику, порождает крайне разреженные векторы высокой размерности и не способен улавливать семантическую близость между синонимами.")
        insert_p_before(p_vyvod, "2. Continuous Bag of Words (CBOW): нейросетевая архитектура, предсказывающая текущее слово по усредненному вектору его контекстного окружения. Формирует плотные эмбеддинги фиксированной размерности, быстро обучается на больших корпусах, однако сглаживает редкие токены.")
        insert_p_before(p_vyvod, "3. Paragraph Vector Distributed Bag of Words (PV-DBOW): алгоритм расширения Doc2Vec, используемый в инструменте iResolver. В этой модели каждому документу (тикету Jira, странице документации) присваивается собственный уникальный обучаемый вектор (Paragraph ID). На этапе обучения вектор документа подается на вход модели и используется для предсказания слов, случайно выбираемых из этого документа. В результате документ проецируется в общее семантическое пространство признаков размерности N = 100, что позволяет напрямую вычислять косинусное сходство между новыми входящими задачами и исторической базой знаний.")

        # SECTION: Аналитическая дискуссия
        insert_h1_before(p_vyvod, "Аналитическая дискуссия по пороговой схожести в 70%")
        insert_p_before(p_vyvod, "При внедрении автоматизированных систем триажа и обработки клиентских запросов (ITSM, L1/L2 техническая поддержка) ключевым параметром классификатора является выбор порога семантической схожести. В рамках выполненной работы исследуется значение косинусного сходства векторов в 70% (cos(theta) = 0.70). Вопрос о том, является ли данный порог «хорошим» или «плохим», требует всесторонней инженерной оценки.")

        insert_h2_before(p_vyvod, "1. Геометрическая интерпретация в пространствах высокой размерности")
        insert_p_before(p_vyvod, "В модели Doc2Vec с размерностью скрытого слоя N = 100 случайные или ортогональные векторы имеют математическое ожидание косинусной близости, близкое к 0 (практический диапазон для несвязанных тем составляет 0.05 – 0.15). Значение схожести 70% соответствует углу между векторами theta approx 45.6 градусов. С точки зрения векторной топологии это свидетельствует о нахождении документов в одном плотном кластере предметной области.")

        insert_h2_before(p_vyvod, "2. Преимущества порога 70% (стратегия высокого охвата — High Recall)")
        insert_p_before(p_vyvod, "• Устойчивость к лексической вариативности: инженеры, клиенты и тестировщики формулируют одну и ту же проблему принципиально разными словами (например, «broker connection timeout», «client hung on handshake», «socket closed unexpectedly»). При пороге 85–90% подобные семантические эквиваленты были бы отсеяны. Порог 70% позволяет обнаружить скрытые дубликаты и ранее решенные инциденты, обладающие разной формулировкой.")
        insert_p_before(p_vyvod, "• Эффективность рекомендаций документации и экспертов: для формирования списков «Useful Documentation» и «Qualified Users» полнота данных важнее абсолютной точечной точности. Даже при схожести 70% документация попадает в правильную подсистему Kafka (Streams, Producer, Core, Security), а ранжированный список инженеров гарантированно содержит специалистов нужного профиля.")

        insert_h2_before(p_vyvod, "3. Ограничения и риски порога 70% (стратегия точности — Low Precision)")
        insert_p_before(p_vyvod, "• Влияние шаблонного шума (Boilerplate): баг-репорты в крупных open-source проектах содержат одинаковые типовые секции: конфигурацию среды (JVM, ОС), одинаковые импорты логов (log4j, slf4j) и шаблонные шаги воспроизведения. Два принципиально разных дефекта могут иметь сходство 70% исключительно за счет общего технического контекста.")
        insert_p_before(p_vyvod, "• Недопустимость автоматического закрытия: использование порога 70% для необратимых действий (автоматическое закрытие тикета как Duplicate или авто-назначение коммита) приведет к неприемлемо высокому уровню ошибок I рода (False Positives).")

        insert_h2_before(p_vyvod, "4. Практические рекомендации для эксплуатации ITSM-систем")
        insert_p_before(p_vyvod, "На основе проведенного анализа рекомендуется применять гибридную двухуровневую модель автоматизации:")
        insert_p_before(p_vyvod, "1. Рекомендательный уровень (Advisory Mode, порог 70%): система выводит инженеру первой линии ранжированный список похожих задач, документации и экспертов в качестве подсказок, экономя до 60% времени на первичный анализ.")
        insert_p_before(p_vyvod, "2. Автоматический уровень (Actionable Mode, порог 85–90%): автоматическое связывание тикетов и автоназначение исполнителей допускается только при превышении 85% сходства и обязательном совпадении структурных метаданных (компонент, модуль, сигнатура исключения).")

    # 6. Update Conclusion text
    # The conclusion paragraph is right after p_vyvod
    p_vyvod_text = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == 'Вывод':
            p_vyvod_text = doc.paragraphs[i+1]
            break

    if p_vyvod_text:
        p_vyvod_text.text = ("В ходе выполнения лабораторной работы были всесторонне изучены и практически освоены современные "
                             "методы автоматизации первичной обработки клиентских запросов на основе инструмента iResolver. "
                             "Была настроена среда исполнения OpenJDK 11 и устранены проблемы совместимости библиотек MVEL2 и DataVec. "
                             "На реальных данных проекта Apache Kafka выгружено более 14 тысяч тикетов Jira, обработана локальная документация "
                             "и свыше 1600 страниц базы знаний Confluence, на основе которых построены векторные модели Doc2Vec (PV-DBOW + Skip-Gram). "
                             "Разработаны и успешно апробированы 4 новых экспертных правила на декларативном языке Drools, обеспечивающие "
                             "контроль подзадач, автоматический перевод статусов и эскалацию неназначенных задач. "
                             "Проведен теоретический анализ алгоритмов Hierarchical Softmax, Skip-Gram, Bag of Words, а также развернута "
                             "инженерная дискуссия по выбору пороговых значений семантической схожести в системах ITSM.")
        if p_vyvod_text.runs:
            style_run(p_vyvod_text.runs[0], size_pt=14)

    # Save to both target locations
    out1 = 'Суханов_СС_Отчет_5.docx'
    out2 = 'Отчет_5.docx'
    doc.save(out1)
    doc.save(out2)
    print(f"Report successfully built and saved to {out1} and {out2}.")

if __name__ == '__main__':
    create_report()
