# -*- coding: utf-8 -*-
"""
Script to replace all 17 figures in Суханов_СС_Отчет_5.docx and Отчет_5.docx
with the user's actual screenshots from the 'скриншоты' directory.
"""

import os
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

def replace_images():
    doc_paths = ['Суханов_СС_Отчет_5.docx', 'Отчет_5.docx']
    folder = 'скриншоты'
    user_images = sorted([f for f in os.listdir(folder) if f.lower().endswith(('.png', '.jpg', '.jpeg'))])
    
    print(f"Found {len(user_images)} screenshots in '{folder}'")
    for idx, f in enumerate(user_images, 1):
        print(f"  Fig 1.{idx} -> {f}")

    if len(user_images) != 17:
        print(f"Warning: Expected 17 images, found {len(user_images)}")

    for doc_path in doc_paths:
        print(f"\nProcessing {doc_path}...")
        doc = docx.Document(doc_path)
        
        # Find all body caption paragraphs (Рис. 1.X) that have an image paragraph preceding them
        figure_slots = []
        for i, p in enumerate(doc.paragraphs):
            txt = p.text.strip()
            # We skip the table of figures at the beginning (P45-P61) by checking if i > 65
            if i > 65 and txt.startswith('Рис. 1.'):
                figure_slots.append((i, txt))

        print(f"Found {len(figure_slots)} figure slots in document.")

        for idx, (cap_idx, cap_text) in enumerate(figure_slots):
            if idx >= len(user_images):
                break
            
            img_file = user_images[idx]
            img_path = os.path.join(folder, img_file)
            
            # The paragraph immediately preceding the caption contains the image
            img_p_idx = cap_idx - 1
            p_img = doc.paragraphs[img_p_idx]
            
            # Clear old content
            p_img.text = ''
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.first_line_indent = Cm(0)
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(3)
            
            # Calculate optimal width
            with Image.open(img_path) as im:
                w, h = im.size
                ratio = w / float(h)
            
            max_width = 16.0
            max_height = 11.5
            
            calc_width = max_width
            calc_height = calc_width / ratio
            if calc_height > max_height:
                calc_width = max_height * ratio
                
            calc_width = min(max_width, max(8.0, calc_width))
            
            run = p_img.add_run()
            run.add_picture(img_path, width=Cm(calc_width))
            print(f"  Replaced Fig 1.{idx+1} (P{img_p_idx}) with {img_file} (width={calc_width:.1f}cm)")

        doc.save(doc_path)
        print(f"Saved updated document: {doc_path}")

if __name__ == '__main__':
    replace_images()
