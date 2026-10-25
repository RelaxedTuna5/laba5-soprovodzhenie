# -*- coding: utf-8 -*-
"""
Generator for Laboratory Report No. 5:
"Автоматизация обработки запросов заказчика"
Student: Суханов Савелий Сергеевич, гр. 5140904/60401
Supervisors: Никифоров И.В., Шишкина В.И.
SPbPU, 2026
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_report():
    doc = docx.Document()
    
    # Page setup (A4, ГОСТ margins: top 2.0cm, bottom 2.0cm, left 3.0cm, right 1.5cm)
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(1.5)
        
    # Styles configuration
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(14)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.3
    normal_style.paragraph_format.space_after = Pt(0)
    normal_style.paragraph_format.space_before = Pt(0)
    normal_style.paragraph_format.first_line_indent = Cm(1.25)
    
    # Helpers
    def add_p(text="", bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=0, indent=1.25, line_spacing=1.3):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.first_line_indent = Cm(indent)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(14)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.italic = True
        r.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_fig(img_path, caption, width_cm=15.0):
        if not os.path.exists(img_path):
            print(f"Warning: Image {img_path} not found!")
            return
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.paragraph_format.first_line_indent = Cm(0)
        r_img = p_img.add_run()
        r_img.add_picture(img_path, width=Cm(width_cm))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(8)
        p_cap.paragraph_format.first_line_indent = Cm(0)
        r_cap = p_cap.add_run(caption)
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(12)
        r_cap.font.italic = False
        r_cap.font.bold = False
        r_cap.font.color.rgb = RGBColor(40, 40, 40)

    def add_code(code_str, caption=""):
        if caption:
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_cap.paragraph_format.space_before = Pt(8)
            p_cap.paragraph_format.space_after = Pt(2)
            p_cap.paragraph_format.first_line_indent = Cm(1.25)
            r_c = p_cap.add_run(caption)
            r_c.font.name = 'Times New Roman'
            r_c.font.size = Pt(12)
            r_c.font.italic = True
            r_c.font.bold = True
            
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Cm(16.5)
        
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8F9FA"/>')
        tcPr.append(shd)
        
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
            f'<w:left w:val="single" w:sz="18" w:space="0" w:color="0969DA"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(code_str)
        r.font.name = 'Consolas'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(36, 41, 46)
        
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(4)
        p_sp.paragraph_format.first_line_indent = Cm(0)

    print("Generator core initialized.")
    return doc, add_p, add_h1, add_h2, add_h3, add_fig, add_code

if __name__ == '__main__':
    create_report()
